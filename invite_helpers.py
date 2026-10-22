"""Helpers for app.py: message rendering, WhatsApp preview, Excel and image handling.

Pure functions only (no Streamlit), so they can be checked from a plain Python shell.
"""
import html
import io
import json
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageOps

# Excel headers are on row 2 (pandas reads with header=1); guests start on row 3.
HEADER_ROW = 2


def render_message(template, name):
    """Fill in {name}. Legacy Markdown bold (**x**) becomes WhatsApp bold (*x*)."""
    # str.replace rather than str.format, so stray { } in the template are harmless.
    return template.replace("**", "*").replace("{name}", str(name))


# --- WhatsApp-style preview -------------------------------------------------

_CODE_BLOCK = re.compile(r"```(.+?)```", re.S)
_INLINE_CODE = re.compile(r"`([^`\n]+)`")


def _style_pattern(marker):
    m = re.escape(marker)
    # WhatsApp only formats when the markers hug the text (no space just inside
    # them), stay on one line and aren't glued to letters/digits outside,
    # e.g. "a*b*c" and "5 * 3 * 2" stay literal.
    return re.compile(
        rf"(?<![^\W_]|{m}){m}(?=[^\s{m}])([^{m}\n]*?[^\s{m}]){m}(?![^\W_]|{m})"
    )


_STYLES = [(_style_pattern("*"), "b"), (_style_pattern("_"), "i"), (_style_pattern("~"), "s")]


def whatsapp_to_html(text, dark=False):
    """Render WhatsApp formatting (*bold* _italic_ ~strike~ ```mono``` `code` > quote)
    as an HTML chat bubble, so the preview shows what the guest will see."""
    if dark:
        bubble, fg, shade = "#005c4b", "#e9edef", "rgba(255,255,255,.12)"
    else:
        bubble, fg, shade = "#d9fdd3", "#111b21", "rgba(0,0,0,.07)"
    code_css = (
        f"font-family:ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;"
        f"background:{shade};border-radius:3px;padding:0 3px"
    )
    quote_css = f"border-left:3px solid {shade};padding-left:8px;opacity:.85"

    s = html.escape(text, quote=False)

    # Park code first so its contents are not formatted.
    stash = []

    def park(match):
        stash.append(f'<code style="{code_css}">{match.group(1)}</code>')
        return f"\x00{len(stash) - 1}\x00"

    s = _CODE_BLOCK.sub(park, s)
    s = _INLINE_CODE.sub(park, s)
    for pattern, tag in _STYLES:
        s = pattern.sub(rf"<{tag}>\1</{tag}>", s)
    s = "\n".join(
        f'<span style="{quote_css}">{line[5:]}</span>' if line.startswith("&gt; ") else line
        for line in s.split("\n")
    )
    s = re.sub(r"\x00(\d+)\x00", lambda m: stash[int(m.group(1))], s)

    return (
        f'<div style="background:{bubble};color:{fg};border-radius:8px;padding:8px 12px;'
        "white-space:pre-wrap;overflow-wrap:anywhere;font-size:15px;line-height:1.45;"
        "font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Helvetica,Arial,sans-serif;"
        f'box-shadow:0 1px 1px rgba(0,0,0,.13)">{s}</div>'
    )


# --- Excel ------------------------------------------------------------------

# Same layout as the sample list; the app only reads "Addressable Names",
# "Add to Invite List?" and "Invited?" - the rest are for the host's planning.
TEMPLATE_COLUMNS = [
    ("Name", 22),
    ("Category", 14),
    ("Addressable Names", 24),
    ("Estimated Count", 15),
    ("Adults", 9),
    ("Kids (Below 8)", 15),
    ("Add to Invite List?", 19),
    ("Invited?", 11),
    ("Notes", 30),
]

TEMPLATE_HELP = [
    "How to fill in this guest list",
    "",
    "1. Add one guest (or family) per row, starting at row 3. Keep the headers in row 2 exactly as they are.",
    "2. Addressable Names (required): how the invite greets them, e.g. \"Ravi & Priya\". It replaces {name} in the message.",
    "3. Add to Invite List? (required): choose Yes to include the guest.",
    "4. Invited? (optional): leave blank - the app sets it to Yes when you click \"Mark sent\".",
    "5. The other columns (Name, Category, counts, Notes) are for your own planning; the app ignores them.",
    "6. Keep \"Invite List\" as the first sheet, save as .xlsx and upload it in the app.",
]


def build_guest_list_template():
    """Return a blank guest-list workbook (.xlsx bytes) in the layout the app expects."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invite List"

    ws["A1"] = "Add guests from row 3 · keep the headers in row 2"
    ws["A1"].font = Font(italic=True, color="666666")

    header_fill = PatternFill("solid", fgColor="D9EAD3")
    letters = {}
    for col, (title, width) in enumerate(TEMPLATE_COLUMNS, start=1):
        letter = get_column_letter(col)
        letters[title] = letter
        cell = ws.cell(row=HEADER_ROW, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        ws.column_dimensions[letter].width = width

    # Row 1 totals, like the sample list.
    for title in ("Estimated Count", "Adults", "Kids (Below 8)"):
        c = letters[title]
        ws[f"{c}1"] = f"=SUBTOTAL(9,{c}3:{c}1000)"
    for title in ("Add to Invite List?", "Invited?"):
        c = letters[title]
        ws[f"{c}1"] = f'=COUNTIF({c}3:{c}1000,"Yes")'

    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(yes_no)
    for title in ("Add to Invite List?", "Invited?"):
        c = letters[title]
        yes_no.add(f"{c}3:{c}1000")

    ws.freeze_panes = "A3"

    help_ws = wb.create_sheet("How to use")
    for i, line in enumerate(TEMPLATE_HELP, start=1):
        help_ws.cell(row=i, column=1, value=line)
    help_ws["A1"].font = Font(bold=True, size=13)
    help_ws.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mark_row(workbook_bytes, excel_row, expected_name, value):
    """Set the 'Invited?' cell of one guest row and return the updated workbook bytes.

    Guests are addressed by Excel row (not by name) so duplicate names are safe;
    the name in that row is checked first as a guard against a shifted sheet.
    """
    wb = load_workbook(io.BytesIO(workbook_bytes))
    # pandas reads the first sheet, so update that one (not wb.active).
    ws = wb.worksheets[0]

    name_col = invited_col = None
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=HEADER_ROW, column=col).value
        if header is None:
            continue
        clean = str(header).strip().lower()
        if clean == "addressable names":
            name_col = col
        elif clean in ("invited", "invited?"):
            invited_col = col

    if name_col is None:
        raise ValueError("Could not find the 'Addressable Names' column in row 2.")

    actual = ws.cell(row=excel_row, column=name_col).value
    if str(actual).strip() != str(expected_name).strip():
        raise ValueError(f"Row {excel_row} holds '{actual}', expected '{expected_name}'.")

    if invited_col is None:
        invited_col = ws.max_column + 1
        ws.cell(row=HEADER_ROW, column=invited_col, value="Invited?")

    ws.cell(row=excel_row, column=invited_col).value = value

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Image ------------------------------------------------------------------

def image_mime(image_bytes):
    """MIME type from the file contents (the extension may lie, e.g. a JPEG named .png)."""
    with Image.open(io.BytesIO(image_bytes)) as im:
        return Image.MIME.get(im.format, "application/octet-stream")


def prepare_image(image_bytes, max_side=2048):
    """Return the image as PNG bytes - the only image type browsers can put on the
    clipboard. Applies phone-camera rotation and caps the size."""
    with Image.open(io.BytesIO(image_bytes)) as im:
        if im.format == "PNG" and max(im.size) <= max_side and im.getexif().get(0x0112, 1) == 1:
            return image_bytes
        im = ImageOps.exif_transpose(im)
        if im.mode not in ("RGB", "RGBA"):
            has_alpha = "A" in im.getbands() or "transparency" in im.info
            im = im.convert("RGBA" if has_alpha else "RGB")
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, format="PNG", optimize=True)
        return buf.getvalue()


# --- Copy buttons -----------------------------------------------------------

_COPY_BUTTONS_HTML = """<!doctype html>
<html><head><meta charset="utf-8"><style>
:root { __COLORS__ }
body { margin: 0; background: transparent;
  font-family: "Source Sans Pro", "Source Sans 3", -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; }
.row { display: flex; flex-wrap: wrap; gap: 8px; }
button { flex: 1 1 150px; padding: 8px 12px; font: inherit; font-size: 15px; cursor: pointer;
  color: var(--fg); background: var(--bg); border: 1px solid var(--border); border-radius: 8px; }
button:hover { border-color: #ff4b4b; color: #ff4b4b; }
#status { margin-top: 6px; font-size: 13px; min-height: 18px; color: var(--muted); }
#status.ok { color: var(--ok); }
#status.err { color: var(--err); }
</style></head><body>
<div class="row">__IMAGE_BUTTON__<button id="copy-text">📋 Copy text</button></div>
<div id="status"></div>
<script>
const data = __PAYLOAD__;
const statusEl = document.getElementById("status");

function show(message, ok) {
  statusEl.textContent = message;
  statusEl.className = ok ? "ok" : "err";
  console.info("[invite-copy] " + message);
}

document.getElementById("copy-text").addEventListener("click", async () => {
  try {
    await navigator.clipboard.writeText(data.text);
  } catch (err) {
    // Fallback for browsers that block the async Clipboard API here.
    const area = document.createElement("textarea");
    area.value = data.text;
    area.style.position = "fixed";
    area.style.opacity = "0";
    document.body.appendChild(area);
    area.select();
    const copied = document.execCommand("copy");
    area.remove();
    if (!copied) {
      show("Couldn't copy the text here - open 'Raw text' below and copy it manually.", false);
      return;
    }
  }
  show("Text copied ✓ - paste it into the chat (⌘V / Ctrl+V) and send.", true);
});

const imageButton = document.getElementById("copy-image");
if (imageButton) {
  // Decode once up front so the click only has to write the ready-made blob.
  const bytes = Uint8Array.from(atob(data.png), (c) => c.charCodeAt(0));
  const blob = new Blob([bytes], { type: "image/png" });
  imageButton.addEventListener("click", async () => {
    try {
      await navigator.clipboard.write([new ClipboardItem({ "image/png": blob })]);
      show("Image copied ✓ - paste it into the chat (⌘V / Ctrl+V), send it, then copy the text.", true);
    } catch (err) {
      show("Couldn't copy the image in this browser (" + err.name + ") - use 'Download Image' and drag it into WhatsApp.", false);
    }
  });
}
</script>
</body></html>
"""


def copy_buttons_html(text, png_base64=None, dark=False):
    """HTML for st.components.v1.html: one-click buttons that put the invite image
    and/or the message text on the clipboard via the browser Clipboard API.

    Works locally and on Streamlit Cloud (the component iframe is allowed
    clipboard-write). The clipboard holds one item at a time, so image and
    text are copied - and pasted - one after the other.
    """
    if dark:
        colors = ("--bg: #262730; --fg: #fafafa; --border: rgba(250,250,250,.2); "
                  "--muted: rgba(250,250,250,.6); --ok: #3dd56d; --err: #ff6c6c;")
    else:
        colors = ("--bg: #ffffff; --fg: #31333f; --border: rgba(49,51,63,.2); "
                  "--muted: rgba(49,51,63,.6); --ok: #177233; --err: #c62828;")
    image_button = '<button id="copy-image">📷 Copy image</button>' if png_base64 else ""
    # Escape "<" so message text can never close the <script> tag early.
    payload = json.dumps({"text": text, "png": png_base64}).replace("<", "\\u003c")
    return (
        _COPY_BUTTONS_HTML.replace("__COLORS__", colors)
        .replace("__IMAGE_BUTTON__", image_button)
        .replace("__PAYLOAD__", payload)
    )
